"""Build procurement/NV_upgrade_shopping_list.xlsx (5 k EUR upgrade plan)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TL = "https://www.thorlabs.com/thorproduct.cfm?partnumber="
DK = "https://www.digikey.es/es/products/detail/mini-circuits/"

# (group, item, part, supplier, qty, unit_price, url, purpose)
TIER_A = [
    ("1. Fluorescence collection & detection",
     "Si switchable-gain photodetector, 75.4 mm² active area", "PDA100A2", "Thorlabs", 1, 399.87,
     TL + "PDA100A2",
     "Dedicated high-collection ensemble channel: 94x larger area than the PDA10A2 and 8 selectable gains, so bandwidth we do not need is traded for transimpedance we do. Directly attacks the 0.1 mV lock-in signal."),
    ("1. Fluorescence collection & detection",
     "Aspheric condenser, Ø25 mm, f = 16 mm, NA = 0.79, AR 650-1050 nm", "ACL25416U-B", "Thorlabs", 2, 31.21,
     TL + "ACL25416U-B",
     "High-NA collection of NV emission close to the diamond, feeding the large-area detector. Far more efficient than re-using the 40x objective for ensemble work."),
    ("1. Fluorescence collection & detection",
     "SM1 lens tube, 1 in", "SM1L10", "Thorlabs", 2, 14.99, TL + "SM1L10",
     "Builds the condenser + filter + detector head as one light-tight tube."),
    ("1. Fluorescence collection & detection",
     "Lens mount for Ø1 in optics, M4 tap", "LMR1/M", "Thorlabs", 2, 16.02, TL + "LMR1/M",
     "Mounting for the condenser lenses."),
    ("1. Fluorescence collection & detection",
     "Ø25 mm longpass filter, 650 nm cut-on", "FELH0650", "Thorlabs", 1, 153.86, TL + "FELH0650",
     "Rejects residual 532 nm and most room light while passing the NV phonon sideband (650-800 nm). Raises contrast by cutting background, not signal."),

    ("2. Stray-light control",
     "Black hardboard, 610 x 610 mm, 3 sheets per pack", "TB4", "Thorlabs", 2, 73.95, TL + "TB4",
     "Rigid light-tight enclosure around the confocal head - the blocker for any low-light / single-NV work."),
    ("2. Stray-light control",
     "Blackout fabric, 1.5 m x 2.7 m", "BK5", "Thorlabs", 1, 61.64, TL + "BK5",
     "Drapes the enclosure and the table edges."),
    ("2. Stray-light control",
     "Matte black aluminium foil, 305 mm x 15.2 m", "BKF12", "Thorlabs", 2, 35.37, TL + "BKF12",
     "Mouldable foil to seal exactly the failure mode we have: cable holes and small apertures."),
    ("2. Stray-light control",
     "Beam trap, 400 nm - 2.5 µm, 30 W", "BT610/M", "Thorlabs", 1, 342.00, TL + "BT610/M",
     "Terminates the transmitted/reflected green safely instead of scattering it around the enclosure. Laser-safety item as well."),

    ("3. Fibre coupling",
     "Achromatic FiberPort, FC/PC & FC/APC, f = 7.5 mm, 400-700 nm", "PAF2-A7A", "Thorlabs", 1, 681.22,
     TL + "PAF2-A7A",
     "Integrated 5-axis fibre launch (x, y, focus + 2 tilts). Our current kinematic-mount approach cannot translate the fibre tip, which is why coupling has never succeeded. NOTE: check the collimated beam diameter first and pick the matching focal length (A4A 0.65 mm / A7A 1.23 mm / A10A / A15A) - all cost the same."),

    ("4. Microwave power characterisation",
     "Fixed attenuator 30 dB, 20 W, DC-18 GHz, SMA", "BW-S30W20+", "DigiKey ES", 1, 203.85,
     DK + "BW-S30W20/16682977",
     "Lets the 16 W amplifier output be measured safely. Needed to find where the ZHL-16W-43-S+ compresses: with +16 dBm drive it is almost certainly saturated, so 'MW power' set on the SMCV may not be the power at the antenna."),
    ("4. Microwave power characterisation",
     "RF power detector, 10 MHz - 8 GHz, DC log output", "ZX47-40-S+", "DigiKey ES", 1, 149.31,
     DK + "ZX47-40-S/21727781",
     "Calibrated MW power axis for the linewidth-vs-power (Rabi) experiment; DC output readable by the multimeter or the Red Pitaya."),
    ("4. Microwave power characterisation",
     "Absorptive SPDT RF switch, DC-5 GHz, TTL control", "ZASWA-2-50DRA+", "DigiKey ES", 1, 189.93,
     "https://www.digikey.es/es/products/result?keywords=ZASWA-2-50DRA%2B",
     "TTL-gated MW from the Red Pitaya: deep 100% amplitude modulation for lock-in, MW on/off normalisation, and a first step towards pulsed sequences."),
    ("4. Microwave power characterisation",
     "Ultra-flexible SMA cable", "(as quoted)", "Mouser", 4, 20.35, "https://www.mouser.es/c/?q=ultra%20flexible%20sma%20cable",
     "Replaces stiff cables that torque the antenna and change coupling between runs. Price from our own earlier quote."),
    ("4. Microwave power characterisation",
     "SMA edge-launch connectors (antenna build)", "(as quoted)", "Mouser", 8, 7.17, "https://www.mouser.es/c/connectors/rf-coaxial-connectors/",
     "For the resonant loop / Omega antenna PCB - stronger B1 is the cheapest route to better ODMR contrast. Price from our own earlier quote."),

    ("5. Temperature & drift",
     "USB temperature + humidity logger with external probe", "TSP01", "Thorlabs", 1, 158.29, TL + "TSP01",
     "Logs lab temperature alongside every ODMR run so drift in D (-74 kHz/K) can be separated from magnetic drift. Enables the thermometry experiment."),
    ("5. Temperature & drift",
     "Flexible polyimide heater with 10 kΩ thermistor", "HT10K", "Thorlabs", 2, 59.79, TL + "HT10K",
     "Deliberately warm the diamond mount to measure dD/dT and calibrate the setup as a thermometer; also quantifies laser-induced heating."),
    ("5. Temperature & drift",
     "10 kΩ thermistor (spares / extra sensing points)", "TH10K", "Thorlabs", 4, 4.84, TL + "TH10K",
     "Additional temperature sensing points (antenna, amplifier, diamond mount)."),

    ("6. Optics cleaning & consumables",
     "Lens tissues, 5 booklets", "MC-5", "Thorlabs", 1, 11.34, TL + "MC-5",
     "Standard drag-wipe cleaning of objective, filters and fibre ends."),
    ("6. Optics cleaning & consumables",
     "Optical-grade cotton-tipped applicators, 100 pack", "CTA10", "Thorlabs", 1, 4.97, TL + "CTA10",
     "Solvent cleaning of connector ferrules and small optics."),
    ("6. Optics cleaning & consumables",
     "Reagent-grade alcohol", "(as quoted)", "Amazon", 1, 8.39, "https://www.amazon.es/s?k=isopropanol+99.9+reagent+grade",
     "Cleaning solvent. Price from our own earlier quote."),
    ("6. Optics cleaning & consumables",
     "Deionised water", "(as quoted)", "Amazon", 1, 10.30, "https://www.amazon.es/s?k=agua+desionizada",
     "Cleaning solvent. Price from our own earlier quote."),
]

TIER_B = [
    ("7. Balanced detection (laser-noise cancellation)",
     "Si switchable-gain photodetector (reference channel)", "PDA100A2", "Thorlabs", 1, 399.87,
     TL + "PDA100A2",
     "Second identical detector on a pick-off of the excitation beam. Fed differentially into the MFLI it cancels common-mode laser intensity noise - the term the noise-budget script is most likely to find dominant."),
    ("7. Balanced detection (laser-noise cancellation)",
     "Ø1 in UVFS wedged beam sampler, AR 350-700 nm", "BSF10-A", "Thorlabs", 1, 72.51, TL + "BSF10-A",
     "Picks off a few % of the green for the reference detector without disturbing the main beam."),

    ("8. Laser power control",
     "Ø1 in linear polariser, 400-700 nm", "LPVISE100-A", "Thorlabs", 1, 103.08, TL + "LPVISE100-A",
     "With the rotation mount, gives continuous laser power control (Malus law) on the linearly polarised DPSS output - needed to take clean laser-power dependences without changing diode current (which changes noise and mode)."),
    ("8. Laser power control",
     "Rotation mount for Ø1 in optics, M4 tap", "RSP1/M", "Thorlabs", 1, 95.73, TL + "RSP1/M",
     "Rotates the polariser; also serves the half-wave plate later."),
    ("8. Laser power control",
     "Absorptive ND filter OD 1.0, SM1-mounted", "NE10A", "Thorlabs", 1, 58.34, TL + "NE10A",
     "Coarse, repeatable power steps for power-dependence series."),
    ("8. Laser power control",
     "Absorptive ND filter OD 2.0, SM1-mounted", "NE20A", "Thorlabs", 1, 58.34, TL + "NE20A",
     "As above, decade step."),

    ("9. Sample & fibre infrastructure",
     "PM patch cable, PANDA, 488 nm, FC/APC, 2 m", "P3-488PM-FC-2", "Thorlabs", 1, 282.24,
     TL + "P3-488PM-FC-2",
     "Spare/second fibre better matched to 532 nm than the 405 nm PM cable, and insurance against an end face damaged during coupling attempts. Keep connector types consistent (APC-APC) across the launch."),
    ("9. Sample & fibre infrastructure",
     "30 mm cage cube precision kinematic rotation platform", "B4CRP/M", "Thorlabs", 1, 333.91, TL + "B4CRP/M",
     "Rotates the diamond + antenna assembly reproducibly: NV-orientation studies, vector magnetometry, and repeatable magnet geometry. Already on our wish list."),
    ("9. Sample & fibre infrastructure",
     "SM1 lens tube, 1 in (light shielding spares)", "SM1L10", "Thorlabs", 2, 14.99, TL + "SM1L10",
     "Extra shielding tubes for the detection path."),
    ("9. Sample & fibre infrastructure",
     "Aspheric condenser (spare / second channel)", "ACL25416U-B", "Thorlabs", 1, 31.21, TL + "ACL25416U-B",
     "Spare for the balanced/second collection channel."),

    ("10. Field calibration",
     "Calibration coil build: magnet wire, 3D-printed former, precision sense resistor",
     "(build)", "RS / Farnell / in-house", 1, 80.00,
     "https://es.rs-online.com/web/c/cables-hilos/hilo-electrico/hilo-de-cobre-esmaltado/",
     "ESTIMATE. A coil of known geometry driven by the KA3010P or the Red Pitaya applies a computable field (N, r, I) so the nT/sqrt(Hz) sensitivity becomes traceable instead of relative, and gives a known AC signal for the WP3 statistics test."),

    ("11. Environment & contingency",
     "Air purifier (optics dust control)", "(as quoted)", "Amazon", 1, 199.00,
     "https://www.amazon.es/s?k=purificador+de+aire+HEPA",
     "Reduces dust settling on the diamond, objective and open optics in a shared room. Price from our own earlier quote."),
    ("11. Environment & contingency",
     "Cotton applicators + lens tissue (restock)", "CTA10 / MC-5", "Thorlabs", 1, 21.28,
     TL + "CTA10",
     "Consumable restock (2x CTA10 + 1x MC-5)."),
    ("11. Environment & contingency",
     "Contingency: shipping, M4/M6 adapters, low-profile screws, SMA savers",
     "(various)", "various", 1, 234.00, "https://www.thorlabs.com/navigation.cfm?guide_id=51",
     "ESTIMATE. Small mechanical hardware identified during assembly plus carriage; keeps the tranche from stalling on 10 EUR parts."),
]

REJECTED = [
    ("Single-photon counting module (e.g. Excelitas SPCM-AQRH)", "~6,000-8,000",
     "The only real route to single-NV work: the APD430A is an analogue detector and a single NV delivers femtoamps. Out of scope for this budget; propose as a separate capital request."),
    ("Zero-order half-wave plate WPH10M-532 (605.44) + mount", "605.44",
     "Would add NV-orientation-selective excitation on top of power control. Deferred - the polariser alone already gives continuous power control. First item to add if funds free up."),
    ("Thorlabs NEK01 ND filter kit (10 filters)", "617.07",
     "Two individual ND filters cover the needed range at 1/5 of the cost."),
    ("3B Scientific Helmholtz coil pair U8481500", "~700",
     "Turnkey but 300 mm coils are oversized for the table and the field is easily computed from a wound coil at ~1/8 the price."),
    ("Second NV diamond with lower NV density", "~1,000-3,700",
     "Would give narrower lines and a resolvable 14N triplet for the Rabi work. Too expensive here; the existing ELSC20 should be tried first."),
    ("Zurich Instruments MFLI options (MF-MD multi-demodulator, MF-PID)", "quote required",
     "The only software licences worth considering: MF-MD would demodulate 1f and 2f simultaneously (derivative + absorption in one pass), MF-PID would move the WP3 feedback loop into hardware. Both need a quote from ZI - request before committing budget."),
    ("Analysis software licence (Origin, MATLAB, LabVIEW)", "n/a",
     "Not recommended: the acquisition and analysis stack is already Python/numpy/matplotlib and is written and tested."),
]

HDR = ["Tier", "Group", "Item", "Part no.", "Supplier", "Qty",
       "Unit price (EUR, ex-VAT)", "Line total (EUR)", "URL", "Purpose / justification"]

wb = Workbook()
ws = wb.active
ws.title = "Shopping list"

F = "Arial"
title_font = Font(name=F, bold=True, size=14)
hdr_font = Font(name=F, bold=True, size=10, color="FFFFFF")
hdr_fill = PatternFill("solid", start_color="1F3864")
grp_font = Font(name=F, bold=True, size=10, color="1F3864")
grp_fill = PatternFill("solid", start_color="D9E2F3")
body = Font(name=F, size=10)
link_font = Font(name=F, size=10, color="0563C1", underline="single")
money_font = Font(name=F, size=10)
input_font = Font(name=F, size=10, color="0000FF")
sub_font = Font(name=F, bold=True, size=10)
sub_fill = PatternFill("solid", start_color="FFF2CC")
tot_fill = PatternFill("solid", start_color="C6E0B4")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(vertical="top", wrap_text=True)

ws["A1"] = "BudgetODMR - NV magnetometer upgrade, procurement plan"
ws["A1"].font = title_font
ws["A2"] = ("All prices EUR, EXCLUDING VAT (21 % ES), as displayed by the vendor on 10 Sep 2026. "
            "Blue = price taken from our own earlier quote or an estimate; black = verified on the linked vendor page.")
ws["A2"].font = Font(name=F, size=9, italic=True)
ws.merge_cells("A1:J1")
ws.merge_cells("A2:J2")

r = 4
for c, h in enumerate(HDR, start=1):
    cell = ws.cell(row=r, column=c, value=h)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = box
ws.freeze_panes = "A5"
r += 1


def write_block(tier_label, rows, start_row):
    rr = start_row
    first_data = None
    last_group = None
    for grp, item, part, sup, qty, price, url, purpose in rows:
        if grp != last_group:
            ws.cell(row=rr, column=1, value=grp).font = grp_font
            for c in range(1, 11):
                ws.cell(row=rr, column=c).fill = grp_fill
                ws.cell(row=rr, column=c).border = box
            ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=10)
            last_group = grp
            rr += 1
        if first_data is None:
            first_data = rr
        estimated = "ESTIMATE" in purpose or "our own earlier quote" in purpose
        vals = [tier_label, grp, item, part, sup, qty, price, None, url, purpose]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=rr, column=c, value=v)
            cell.font = body
            cell.alignment = wrap
            cell.border = box
        ws.cell(row=rr, column=7).font = input_font if estimated else money_font
        ws.cell(row=rr, column=7).number_format = '#,##0.00 "EUR"'
        ws.cell(row=rr, column=8, value=f"=F{rr}*G{rr}")
        ws.cell(row=rr, column=8).number_format = '#,##0.00 "EUR"'
        ws.cell(row=rr, column=8).font = money_font
        ws.cell(row=rr, column=8).border = box
        lk = ws.cell(row=rr, column=9)
        lk.hyperlink = url
        lk.value = url
        lk.font = link_font
        rr += 1
    return first_data, rr - 1, rr


a_first, a_last, r = write_block("A - core (3 k tranche)", TIER_A, r)
ws.cell(row=r, column=3, value="TIER A SUBTOTAL - core upgrade (3,000 EUR tranche)").font = sub_font
ws.cell(row=r, column=8, value=f"=SUM(H{a_first}:H{a_last})")
for c in range(1, 11):
    ws.cell(row=r, column=c).fill = sub_fill
    ws.cell(row=r, column=c).border = box
ws.cell(row=r, column=8).number_format = '#,##0.00 "EUR"'
ws.cell(row=r, column=8).font = sub_font
a_sub = r
r += 2

b_first, b_last, r = write_block("B - extension (2 k tranche)", TIER_B, r)
ws.cell(row=r, column=3, value="TIER B SUBTOTAL - extension (2,000 EUR tranche)").font = sub_font
ws.cell(row=r, column=8, value=f"=SUM(H{b_first}:H{b_last})")
for c in range(1, 11):
    ws.cell(row=r, column=c).fill = sub_fill
    ws.cell(row=r, column=c).border = box
ws.cell(row=r, column=8).number_format = '#,##0.00 "EUR"'
ws.cell(row=r, column=8).font = sub_font
b_sub = r
r += 1

ws.cell(row=r, column=3, value="GRAND TOTAL (ex-VAT)").font = Font(name=F, bold=True, size=11)
ws.cell(row=r, column=8, value=f"=H{a_sub}+H{b_sub}")
for c in range(1, 11):
    ws.cell(row=r, column=c).fill = tot_fill
    ws.cell(row=r, column=c).border = box
ws.cell(row=r, column=8).number_format = '#,##0.00 "EUR"'
ws.cell(row=r, column=8).font = Font(name=F, bold=True, size=11)
gt = r
r += 1
ws.cell(row=r, column=3, value="Memo: same total including 21 % Spanish VAT").font = Font(name=F, italic=True, size=9)
ws.cell(row=r, column=8, value=f"=H{gt}*1.21")
ws.cell(row=r, column=8).number_format = '#,##0.00 "EUR"'
ws.cell(row=r, column=8).font = Font(name=F, italic=True, size=9)
ws.cell(row=r, column=9, value="Check with the finance office which basis the 3 k + 2 k awards are quoted on.")
ws.cell(row=r, column=9).font = Font(name=F, italic=True, size=9)
r += 2

ws.cell(row=r, column=1, value="Budget check").font = Font(name=F, bold=True, size=11)
r += 1
for label, formula in (("Tier A vs 3,000 EUR available", f"=3000-H{a_sub}"),
                       ("Tier B vs 2,000 EUR available", f"=2000-H{b_sub}"),
                       ("Total headroom vs 5,000 EUR", f"=5000-H{gt}")):
    ws.cell(row=r, column=3, value=label).font = body
    ws.cell(row=r, column=8, value=formula)
    ws.cell(row=r, column=8).number_format = '#,##0.00 "EUR"'
    ws.cell(row=r, column=8).font = body
    r += 1

widths = {"A": 22, "B": 30, "C": 46, "D": 17, "E": 15, "F": 6,
          "G": 15, "H": 14, "I": 52, "J": 82}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=5, max_row=gt):
    ws.row_dimensions[row[0].row].height = None

# ---------------- second sheet: considered & rejected -----------------------
ws2 = wb.create_sheet("Considered - not bought")
ws2["A1"] = "Items evaluated and deliberately excluded"
ws2["A1"].font = title_font
ws2.merge_cells("A1:C1")
for c, h in enumerate(["Item", "Indicative price (EUR)", "Why it is not in this request"], start=1):
    cell = ws2.cell(row=3, column=c, value=h)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.border = box
    cell.alignment = Alignment(vertical="center", wrap_text=True)
rr = 4
for item, price, why in REJECTED:
    for c, v in enumerate([item, price, why], start=1):
        cell = ws2.cell(row=rr, column=c, value=v)
        cell.font = body
        cell.alignment = wrap
        cell.border = box
    rr += 1
ws2.column_dimensions["A"].width = 54
ws2.column_dimensions["B"].width = 22
ws2.column_dimensions["C"].width = 92

wb.save("NV_upgrade_shopping_list.xlsx")
print("written NV_upgrade_shopping_list.xlsx")
